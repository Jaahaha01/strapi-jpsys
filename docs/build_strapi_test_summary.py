from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
RESULTS_DOCS = RESULTS / "docs"
RESULTS_K6 = RESULTS / "k6"
OUT_DOCX = RESULTS_DOCS / "strapi-test-summary.docx"
OUT_MD = RESULTS_DOCS / "strapi-test-summary.md"

RESULTS_DOCS.mkdir(parents=True, exist_ok=True)


def set_run_font(run, *, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
            row.cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def style_table(table, widths, header_fill="E8EEF5"):
    set_table_width(table, widths)
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_run_font(run, size=10.5, color="000000")
    for cell in table.rows[0].cells:
        shade_cell(cell, header_fill)
        for p in cell.paragraphs:
            for run in p.runs:
                set_run_font(run, size=10.5, color="000000", bold=True)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=15, color="2E74B5", bold=True)
    else:
        set_run_font(run, size=12.5, color="1F4D78", bold=True)
    return p


def add_body(doc, text, *, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=11, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, size=11)
    else:
        r = p.add_run(text)
        set_run_font(r, size=11)
    return p


def add_key_value_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(1.45), Inches(5.05)]
    for label, value in rows:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = value
    style_table(table, widths, header_fill="F2F4F7")
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            if j == 0:
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        set_run_font(run, size=10.5, bold=True)
            else:
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        set_run_font(run, size=10.5)
    return table


def read_csv_rows(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def make_doc():
    doc = Document()

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(3)
    r = title.add_run("Strapi E2E Test Summary")
    set_run_font(r, size=22, color="000000", bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(10)
    r = subtitle.add_run("Playwright admin journeys and k6 load test results")
    set_run_font(r, size=11, color="555555")

    add_key_value_table(
        doc,
        [
            ("Project", "C:\\JapanSys\\strapi-jpsys"),
            ("Captured", "May 29, 2026"),
            ("Scope", "Playwright admin journeys for Strapi admin plus k6 load testing at 10, 50, and 100 VUs"),
            ("Outputs", "DOCX summary plus linked CSV/HTML artifacts in the repo"),
        ],
    )

    add_heading(doc, "Executive Summary", level=1)
    add_body(
        doc,
        "The load-test side is clean: all three k6 runs completed with zero request failures, and p95 stayed well below the configured thresholds. The Playwright side is mixed: Chromium passes the four admin journeys, while the full multi-browser run fails on Firefox launch before the app logic can finish running.",
    )

    add_heading(doc, "Playwright E2E", level=1)
    play_table = doc.add_table(rows=1, cols=3)
    play_table.style = "Table Grid"
    play_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    play_headers = ["Run", "Result", "Notes"]
    for idx, header in enumerate(play_headers):
        play_table.rows[0].cells[idx].text = header
    play_rows = [
        [
            "Full suite (Chromium + Firefox + WebKit)",
            "8 passed, 1 failed, 3 did not run",
            "Firefox crashed during headless launch with a graphics compositor error; the failure happened before the journey logic could complete.",
        ],
        [
            "Chromium only",
            "4/4 passed",
            "Login, create, edit, and delete/logout journeys all completed successfully.",
        ],
        [
            "Current test design",
            "Needs rework",
            "The spec still behaves like collection-style CRUD, but the Contact content type is a singleType in this repo.",
        ],
    ]
    for row in play_rows:
        cells = play_table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    style_table(play_table, [Inches(1.55), Inches(1.3), Inches(3.65)], header_fill="E8EEF5")
    for row in play_table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, size=10.2)

    add_heading(doc, "k6 Load Test", level=1)
    k6_rows = read_csv_rows(RESULTS_K6 / "report-overall.csv")
    k6_table = doc.add_table(rows=1, cols=7)
    k6_table.style = "Table Grid"
    k6_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    k6_headers = ["VUs", "Requests", "Req/s", "Avg ms", "p95 ms", "Error %", "Verdict"]
    for idx, header in enumerate(k6_headers):
        k6_table.rows[0].cells[idx].text = header
    for row in k6_rows:
        cells = k6_table.add_row().cells
        cells[0].text = row["vus"]
        cells[1].text = row["total_requests"]
        cells[2].text = row["requests_per_second"]
        cells[3].text = row["avg_ms"]
        cells[4].text = row["p95_ms"]
        cells[5].text = row["custom_error_percent"]
        cells[6].text = row["result"]
    style_table(k6_table, [Inches(0.65), Inches(0.95), Inches(0.8), Inches(0.8), Inches(0.8), Inches(0.8), Inches(0.9)], header_fill="E8EEF5")
    for row in k6_table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_run_font(run, size=10.2)

    add_body(
        doc,
        "The PUT endpoint is consistently the slowest one under load, which is expected because it writes data instead of only reading it. Even at 100 VUs, the run stayed within the configured p95 threshold and reported no failed requests.",
    )

    add_heading(doc, "100 VUs Endpoint Snapshot", level=1)
    endpoint_rows = read_csv_rows(RESULTS_K6 / "report-endpoints.csv")
    endpoint_100 = [row for row in endpoint_rows if row["vus"] == "100"]
    ep_table = doc.add_table(rows=1, cols=8)
    ep_table.style = "Table Grid"
    ep_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    ep_headers = ["Method", "Endpoint", "Avg ms", "p95 ms", "Max ms", "Requests", "Failures", "Failure %"]
    for idx, header in enumerate(ep_headers):
        ep_table.rows[0].cells[idx].text = header
    for row in endpoint_100:
        cells = ep_table.add_row().cells
        cells[0].text = row["method"]
        cells[1].text = row["endpoint"]
        cells[2].text = row["avg_ms"]
        cells[3].text = row["p95_ms"]
        cells[4].text = row["max_ms"]
        cells[5].text = row["requests"]
        cells[6].text = row["failures"]
        cells[7].text = row["failure_percent"]
    style_table(ep_table, [Inches(0.75), Inches(1.8), Inches(0.7), Inches(0.75), Inches(0.75), Inches(0.75), Inches(0.65), Inches(0.7)], header_fill="E8EEF5")
    for row in ep_table.rows:
        for idx, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 1 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_run_font(run, size=10)

    add_heading(doc, "Key Observations", level=1)
    obs_table = doc.add_table(rows=1, cols=2)
    obs_table.style = "Table Grid"
    obs_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    obs_table.rows[0].cells[0].text = "Observation"
    obs_table.rows[0].cells[1].text = "Why it matters"
    observations = [
        (
            "Firefox fails before the app logic runs.",
            "The current multi-browser Playwright result is not a product regression; it looks like a headless graphics/runtime problem on this machine.",
        ),
        (
            "The Contact admin journey is mismatched with the schema.",
            "The test still behaves like collection CRUD, but Contact is a singleType, so the suite should be rewritten against the real admin flow.",
        ),
        (
            "k6 stayed within threshold at every level.",
            "The Strapi API is responsive enough under the tested load, and the repo already has CSV output that is easy to chart in Google Sheets.",
        ),
    ]
    for left, right in observations:
        cells = obs_table.add_row().cells
        cells[0].text = left
        cells[1].text = right
    style_table(obs_table, [Inches(2.05), Inches(4.45)], header_fill="E8EEF5")
    for row in obs_table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, size=10.2)

    add_heading(doc, "Recommended Next Steps", level=1)
    next_steps = [
        "Rewrite the Playwright journeys to follow the real singleType admin flow instead of collection-style create/edit/delete behavior.",
        "Keep Chromium as the required browser in CI until the Firefox launch crash is fixed or isolated.",
        "Use the existing k6 CSV outputs for the charts and keep the stage-mode summary as a separate appendix if you want to compare ramp behavior later.",
    ]
    for text in next_steps:
        add_body(doc, text)

    add_heading(doc, "Artifacts", level=1)
    artifacts = doc.add_table(rows=1, cols=2)
    artifacts.style = "Table Grid"
    artifacts.alignment = WD_TABLE_ALIGNMENT.CENTER
    artifacts.rows[0].cells[0].text = "File"
    artifacts.rows[0].cells[1].text = "Purpose"
    artifact_rows = [
        ("results/e2e/playwright-report/index.html", "HTML report for the Playwright run"),
        ("results/e2e/test-results/", "Trace artifacts for the failing Firefox run"),
        ("results/k6/report-overall.csv", "Overall chart data for 10/50/100 VUs"),
        ("results/k6/report-endpoints.csv", "Endpoint chart data for 100 VUs"),
        ("results/k6/LOAD_TEST_SUMMARY.md", "Markdown summary of the k6 run"),
    ]
    for file_name, purpose in artifact_rows:
        cells = artifacts.add_row().cells
        cells[0].text = file_name
        cells[1].text = purpose
    style_table(artifacts, [Inches(2.35), Inches(4.15)], header_fill="F2F4F7")
    for row in artifacts.rows:
        for idx, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_run_font(run, size=10)

    doc.save(OUT_DOCX)

    md = f"""# Strapi E2E Test Summary

Generated from the latest Playwright and k6 runs in the repo.

## Executive Summary

- k6 passed at 10, 50, and 100 VUs with zero request failures.
- Playwright Chromium passed all four admin journeys.
- Firefox failed during browser launch in headless mode, before the app logic completed.
- The current Playwright journeys still model collection-style CRUD even though Contact is a `singleType`.

## k6 Summary

| VUs | Requests | Req/s | Avg ms | p95 ms | Error % | Verdict |
|---:|---:|---:|---:|---:|---:|:---|
"""
    for row in k6_rows:
        md += f"| {row['vus']} | {row['total_requests']} | {row['requests_per_second']} | {row['avg_ms']} | {row['p95_ms']} | {row['custom_error_percent']} | {row['result']} |\n"

    md += "\n## Playwright Summary\n\n"
    md += "- Full multi-browser run: 8 passed, 1 failed, 3 did not run.\n"
    md += "- Chromium-only rerun: 4/4 passed.\n"
    md += "- Firefox launch error: `RenderCompositorSWGL failed mapping default framebuffer`.\n"
    md += "\n## Next Steps\n\n"
    md += "- Rewrite the admin journeys for the real singleType flow.\n"
    md += "- Keep Chromium required in CI until Firefox is fixed or isolated.\n"
    md += "- Use the CSV outputs for charting and reporting.\n"

    OUT_MD.write_text(md, encoding="utf-8")


if __name__ == "__main__":
    RESULTS_DOCS.mkdir(parents=True, exist_ok=True)
    make_doc()
    print(f"Wrote {OUT_DOCX}")
    print(f"Wrote {OUT_MD}")
